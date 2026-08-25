import csv
import sys
import threading
import tkinter as tk
from datetime import datetime
from pathlib import Path
from tkinter import filedialog, ttk

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

sys.path.insert(0, str(Path(__file__).parent))
from monitor import (
    CONFIG_FILE, build_email_body, fetch_copy,
    fields_match, load_json, send_email,
)

FIELDS = [("title", "Title"), ("meta_description", "Meta Description"), ("intro_copy", "Intro Copy"), ("h1", "H1")]

BG      = "#f5f5f5"
WHITE   = "#ffffff"
DARK    = "#1a1a1a"
MUTED   = "#888888"
GREEN_BG = "#edfaf1";  GREEN_FG = "#1a7a3e"
RED_BG   = "#fdecea";  RED_FG   = "#b71c1c"
BLUE    = "#2563eb"
DIVIDER = "#e0e0e0"
ROW_ALT = "#fafafa"

COL_FIELD_W = 140   # px – label column


class ScrollableFrame(tk.Frame):
    def __init__(self, parent, **kwargs):
        super().__init__(parent, **kwargs)
        self.canvas = tk.Canvas(self, bg=WHITE, highlightthickness=0)
        sb = tk.Scrollbar(self, orient="vertical", command=self.canvas.yview)
        self.inner = tk.Frame(self.canvas, bg=WHITE)
        self.canvas.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self.canvas.pack(side="left", fill="both", expand=True)
        self._win = self.canvas.create_window((0, 0), window=self.inner, anchor="nw")
        self.inner.bind("<Configure>", lambda _: self.canvas.configure(scrollregion=self.canvas.bbox("all")))
        self.canvas.bind("<Configure>", lambda e: self.canvas.itemconfig(self._win, width=e.width))
        self.canvas.bind_all("<MouseWheel>", lambda e: self.canvas.yview_scroll(-1*(e.delta//120), "units"))


class App(tk.Tk):
    def __init__(self):
        super().__init__()
        self.cfg = load_json(CONFIG_FILE)
        self.csv_rows: list[dict] = []
        self.current_saved: dict = self.cfg.get("saved_copy", {})
        self._result_labels: list[tk.Label] = []   # for resize wraplength update

        self.title("Site Monitor")
        self.configure(bg=BG)
        self.minsize(720, 500)
        self.resizable(True, True)
        self._build()
        self.geometry("960x640")
        self.update_idletasks()
        w, h = self.winfo_width(), self.winfo_height()
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")
        self.bind("<Configure>", self._on_resize)

    # ── Build UI ──────────────────────────────────────────────────────────────

    def _build(self):
        # Header
        hdr = tk.Frame(self, bg=DARK, padx=24, pady=20)
        hdr.pack(fill="x")
        tk.Label(hdr, text="Site Monitor", font=("Arial", 10), bg=DARK, fg=MUTED).pack(anchor="w")
        tk.Label(hdr, text="SEO Copy Checker", font=("Arial", 13, "bold"), bg=DARK, fg=WHITE).pack(anchor="w", pady=(2,0))

        # CSV panel
        csv_panel = tk.Frame(self, bg="#ececec", padx=24, pady=12)
        csv_panel.pack(fill="x")

        btn_row = tk.Frame(csv_panel, bg="#ececec")
        btn_row.pack(fill="x")

        tk.Button(btn_row, text="Upload File (CSV or XLSX)",
                  font=("Arial", 10, "bold"), bg=DARK, fg=WHITE,
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#333", activeforeground=WHITE,
                  command=self._load_file).pack(side="left", padx=(0, 8))

        tk.Button(btn_row, text="Download Template (.xlsx)",
                  font=("Arial", 10), bg="#ececec", fg="#555",
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#ddd",
                  command=self._download_template).pack(side="left")

        # URL selector (shown after file load)
        self._selector_frame = tk.Frame(btn_row, bg="#ececec")
        self._selector_frame.pack(side="right")
        tk.Label(self._selector_frame, text="Select URL:", font=("Arial", 10),
                 bg="#ececec", fg="#555").pack(side="left", padx=(0, 6))
        self._csv_url_var = tk.StringVar()
        self._csv_selector = ttk.Combobox(self._selector_frame, textvariable=self._csv_url_var,
                                          state="readonly", width=38, font=("Arial", 10))
        self._csv_selector.pack(side="left")
        self._csv_selector.bind("<<ComboboxSelected>>", self._on_csv_url_select)
        self._selector_frame.pack_forget()

        tk.Label(csv_panel,
                 text="Template format — Col A: URL  |  Col B: Title  |  Col C: Meta Description  |  Col D: Intro Copy  |  Col E: H1  (row 1 = headers, data from row 2)",
                 font=("Arial", 9), bg="#ececec", fg=MUTED).pack(anchor="w", pady=(8, 0))

        # Toolbar
        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")
        toolbar = tk.Frame(self, bg=BG, padx=24, pady=14)
        toolbar.pack(fill="x")
        toolbar.columnconfigure(1, weight=1)

        self._run_btn = tk.Button(toolbar, text="Run Check",
                                  font=("Arial", 11, "bold"), bg=BLUE, fg=WHITE,
                                  relief="flat", padx=14, pady=8, cursor="hand2",
                                  activebackground="#1d4ed8", activeforeground=WHITE,
                                  command=self._run_single)
        self._run_btn.grid(row=0, column=0, padx=(0, 8))

        self._run_all_btn = tk.Button(toolbar, text="Run All",
                                      font=("Arial", 11, "bold"), bg="#16a34a", fg=WHITE,
                                      relief="flat", padx=14, pady=8, cursor="hand2",
                                      activebackground="#15803d", activeforeground=WHITE,
                                      state="disabled", command=self._run_all)
        self._run_all_btn.grid(row=0, column=1, sticky="w", padx=(0, 12))

        self._url_var = tk.StringVar(value=self.cfg.get("url", ""))
        url_entry = tk.Entry(toolbar, textvariable=self._url_var, font=("Arial", 11),
                             relief="flat", bg=WHITE, fg=DARK, insertbackground=DARK,
                             highlightthickness=1, highlightbackground=DIVIDER, highlightcolor=BLUE)
        url_entry.grid(row=0, column=2, sticky="ew", ipady=8)
        url_entry.bind("<Return>", lambda _: self._run_single())
        toolbar.columnconfigure(2, weight=1)

        self._last_checked_var = tk.StringVar(value="Not checked yet")
        tk.Label(toolbar, textvariable=self._last_checked_var, font=("Arial", 10),
                 bg=BG, fg=MUTED).grid(row=1, column=0, columnspan=3, sticky="w", pady=(6,0))

        # Column header bar
        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")
        col_hdr = tk.Frame(self, bg="#f0f0f0", padx=24, pady=8)
        col_hdr.pack(fill="x")
        col_hdr.columnconfigure(0, minsize=COL_FIELD_W, weight=0)
        col_hdr.columnconfigure(1, weight=1)
        col_hdr.columnconfigure(2, weight=1)
        for ci, txt in enumerate(["FIELD", "LIVE COPY", "YOUR SAVED COPY"]):
            tk.Label(col_hdr, text=txt, anchor="w", font=("Arial", 9, "bold"),
                     bg="#f0f0f0", fg=MUTED).grid(row=0, column=ci, sticky="ew", padx=(0,16))

        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")

        # Scrollable results area
        self._scroll = ScrollableFrame(self, bg=WHITE)
        self._scroll.pack(fill="both", expand=True)

        # Status bar
        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")
        self._status_var = tk.StringVar(value="Ready — enter a URL or upload a file to begin")
        sb_frame = tk.Frame(self, bg="#f0f0f0", padx=24, pady=10)
        sb_frame.pack(fill="x")
        tk.Label(sb_frame, textvariable=self._status_var, font=("Arial", 10),
                 bg="#f0f0f0", fg=MUTED, anchor="w").pack(fill="x")

    # ── Results rendering ─────────────────────────────────────────────────────

    def _clear_results(self):
        for w in self._scroll.inner.winfo_children():
            w.destroy()
        self._result_labels.clear()

    def _render_result(self, url: str, live: dict, saved: dict,
                       mismatches: list, error: str = ""):
        inner = self._scroll.inner
        all_match = not mismatches and not error
        icon = "✅" if all_match else ("⚠️" if error else "❌")

        # URL header
        url_bg = "#1e3a1e" if all_match else ("#5a1a1a" if error else "#2d2d2d")
        url_hdr = tk.Frame(inner, bg=url_bg)
        url_hdr.pack(fill="x")
        tk.Label(url_hdr, text=f"{icon}  {url}", font=("Arial", 10, "bold"),
                 bg=url_bg, fg=WHITE, padx=24, pady=10, anchor="w").pack(fill="x")

        if error:
            tk.Label(inner, text=f"Error: {error}", font=("Arial", 10),
                     bg=RED_BG, fg=RED_FG, padx=24, pady=10, anchor="w").pack(fill="x")
        else:
            for ri, (key, label) in enumerate(FIELDS):
                row_bg = WHITE if ri % 2 == 0 else ROW_ALT
                live_val  = live.get(key, "") or "—"
                saved_val = saved.get(key, "")
                is_ph = not saved_val or saved_val.startswith("PASTE THE EXACT")

                if is_ph:
                    cell_bg, cell_fg = row_bg, MUTED
                elif key not in mismatches:
                    cell_bg, cell_fg = GREEN_BG, GREEN_FG
                else:
                    cell_bg, cell_fg = RED_BG, RED_FG

                row_frame = tk.Frame(inner, bg=row_bg)
                row_frame.pack(fill="x")
                row_frame.columnconfigure(0, minsize=COL_FIELD_W, weight=0)
                row_frame.columnconfigure(1, weight=1)
                row_frame.columnconfigure(2, weight=1)

                tk.Label(row_frame, text=label, anchor="nw", font=("Arial", 10, "bold"),
                         bg=row_bg, fg="#444", padx=24, pady=12).grid(row=0, column=0, sticky="nsew")

                lv = tk.Label(row_frame, text=live_val, anchor="nw", font=("Arial", 10),
                              bg=row_bg, fg="#222", padx=8, pady=12, wraplength=280, justify="left")
                lv.grid(row=0, column=1, sticky="nsew", padx=(0, 8))

                sv = tk.Label(row_frame, text=saved_val if not is_ph else "(not set)",
                              anchor="nw", font=("Arial", 10),
                              bg=cell_bg, fg=cell_fg, padx=8, pady=12, wraplength=280, justify="left")
                sv.grid(row=0, column=2, sticky="nsew", padx=(0, 24))

                self._result_labels.extend([lv, sv])
                tk.Frame(inner, bg=DIVIDER, height=1).pack(fill="x")

        tk.Frame(inner, bg="#aaaaaa", height=2).pack(fill="x", pady=(0, 6))

    # ── File loading ──────────────────────────────────────────────────────────

    def _load_file(self):
        path = filedialog.askopenfilename(
            title="Select file",
            filetypes=[("Spreadsheet files", "*.csv *.xlsx"),
                       ("CSV", "*.csv"), ("Excel", "*.xlsx"), ("All", "*.*")],
        )
        if not path:
            return
        try:
            rows = self._read_xlsx(path) if path.lower().endswith(".xlsx") else self._read_csv(path)
        except Exception as e:
            self._status_var.set(f"Failed to read file: {e}")
            return

        self.csv_rows = [r for r in rows if str(r.get("URL", "")).strip()]
        if not self.csv_rows:
            self._status_var.set("No URLs found. Check column A is labelled 'URL'.")
            return

        urls = [r["URL"].strip() for r in self.csv_rows]
        self._csv_selector["values"] = urls
        self._csv_selector.current(0)
        self._selector_frame.pack(side="right")
        self._apply_row(0)
        count = len(self.csv_rows)
        self._run_all_btn.config(state="normal", text=f"Run All ({count})")
        self._status_var.set(f"{count} URL(s) loaded. Select one and press Run Check, or press Run All.")

    def _read_csv(self, path: str) -> list[dict]:
        with open(path, newline="", encoding="utf-8-sig") as f:
            reader = csv.DictReader(f)
            reader.fieldnames = [h.strip() for h in (reader.fieldnames or [])]
            return list(reader)

    def _read_xlsx(self, path: str) -> list[dict]:
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb.active
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            return []
        headers = [str(h).strip() if h is not None else "" for h in rows[0]]
        result = []
        for row in rows[1:]:
            if any(c for c in row if c is not None):
                result.append({headers[i]: (str(v).strip() if v is not None else "") for i, v in enumerate(row)})
        return result

    def _on_csv_url_select(self, _):
        self._apply_row(self._csv_selector.current())

    def _apply_row(self, idx: int):
        if idx < 0 or idx >= len(self.csv_rows):
            return
        row = self.csv_rows[idx]
        self._url_var.set(row.get("URL", "").strip())
        self.current_saved = {
            "title": row.get("Title", "").strip(),
            "meta_description": row.get("Meta Description", "").strip(),
            "intro_copy": row.get("Intro Copy", "").strip(),
            "h1": row.get("H1", "").strip(),
        }

    # ── Template download ─────────────────────────────────────────────────────

    def _download_template(self):
        path = filedialog.asksaveasfilename(
            title="Save template",
            defaultextension=".xlsx",
            initialfile="site-monitor-template.xlsx",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if not path:
            return
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "URLs"

        headers = ["URL", "Title", "Meta Description", "Intro Copy", "H1"]
        col_widths = [45, 55, 70, 80, 55]
        header_fill = PatternFill("solid", fgColor="1A1A1A")
        header_font = Font(bold=True, color="FFFFFF", size=11)

        for ci, (h, w) in enumerate(zip(headers, col_widths), start=1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center", wrap_text=False)
            ws.column_dimensions[cell.column_letter].width = w

        ws.row_dimensions[1].height = 22

        examples = [
            ["https://uk.pandora.net/en/",
             "Official Pandora™ UK | Charms & Bracelets | Women's Jewellery",
             "Find your style in Pandora's jewellery collection. Shop charms, earrings, rings, necklaces and bracelets for the pieces that match your look, today.",
             "Your expected intro paragraph copy",
             "Your expected H1"],
            ["https://example.com/",
             "Your expected page title",
             "Your expected meta description",
             "Your expected intro paragraph copy",
             "Your expected H1"],
        ]
        for ri, ex in enumerate(examples, start=2):
            for ci, val in enumerate(ex, start=1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[ri].height = 30

        wb.save(path)
        self._status_var.set(f"Template saved — open it, fill in your URLs and copy, then upload it here.")

    # ── Check logic ───────────────────────────────────────────────────────────

    def _set_busy(self, busy: bool, label: str = "Run Check"):
        state = "disabled" if busy else "normal"
        self._run_btn.config(state=state, text="Checking…" if busy else label)
        self._run_all_btn.config(state=state if not busy else "disabled")

    def _run_single(self):
        url = self._url_var.get().strip()
        if not url:
            self._status_var.set("Please enter a URL.")
            return
        self._set_busy(True)
        self._status_var.set("Fetching page…")
        self._clear_results()
        threading.Thread(target=self._do_single, args=(url, dict(self.current_saved)), daemon=True).start()

    def _do_single(self, url: str, saved: dict):
        try:
            live = fetch_copy(url)
            _, mismatches = fields_match(live, saved)
            error = ""
        except Exception as e:
            live, mismatches, error = {}, [], str(e)

        all_match = not mismatches and not error
        if all_match:
            try:
                html = build_email_body(url, live, saved, all_match)
                send_email(self.cfg, f"[Site Monitor] ✅ Copy is live — {url}", html)
                email_note = "Email sent."
            except Exception as e:
                email_note = f"Email failed: {e}"
        else:
            email_note = "No email — copy doesn't match yet."

        def update():
            self._render_result(url, live, saved, mismatches, error)
            now = datetime.now().strftime("%d %b %Y, %H:%M:%S")
            self._last_checked_var.set(f"Last checked: {now}")
            icon = "✅" if all_match else "❌"
            self._status_var.set(f"{icon} {'All fields match' if all_match else 'Some fields differ'}  ·  {email_note}")
            self._set_busy(False)

        self.after(0, update)

    def _run_all(self):
        if not self.csv_rows:
            return
        self._run_btn.config(state="disabled")
        self._run_all_btn.config(state="disabled", text="Running…")
        self._status_var.set(f"Starting bulk check of {len(self.csv_rows)} URLs…")
        self._clear_results()
        threading.Thread(target=self._do_bulk, daemon=True).start()

    def _do_bulk(self):
        results = []
        total = len(self.csv_rows)
        for i, row in enumerate(self.csv_rows):
            url = row.get("URL", "").strip()
            saved = {
                "title": row.get("Title", "").strip(),
                "meta_description": row.get("Meta Description", "").strip(),
                "intro_copy": row.get("Intro Copy", "").strip(),
                "h1": row.get("H1", "").strip(),
            }
            self.after(0, self._status_var.set, f"Checking {i+1}/{total}: {url}")
            try:
                live = fetch_copy(url)
                _, mismatches = fields_match(live, saved)
                error = ""
            except Exception as e:
                live, mismatches, error = {}, [], str(e)

            results.append((url, live, saved, mismatches, error))
            self.after(0, self._render_result, url, live, saved, mismatches, error)

        # Send email only if at least one URL matched
        matched_results = [(u, l, s, m, e) for u, l, s, m, e in results if not m and not e]
        matches = len(matched_results)
        if matched_results:
            try:
                html = self._build_bulk_email(matched_results, total)
                subj = f"[Site Monitor] ✅ {matches}/{total} URLs matched"
                send_email(self.cfg, subj, html)
                email_note = f"Email sent — {matches} match{'es' if matches > 1 else ''}."
            except Exception as e:
                email_note = f"Email failed: {e}"
        else:
            email_note = "No email — no URLs matched yet."

        def done():
            matches = sum(1 for _, _, _, m, e in results if not m and not e)
            now = datetime.now().strftime("%d %b %Y, %H:%M:%S")
            self._last_checked_var.set(f"Last checked: {now}")
            self._status_var.set(f"Done — {matches}/{total} matched  ·  {email_note}")
            self._run_btn.config(state="normal")
            self._run_all_btn.config(state="normal", text=f"Run All ({total})")

        self.after(0, done)

    def _build_bulk_email(self, results: list, total: int) -> str:
        now = datetime.now().strftime("%d %b %Y, %H:%M")
        matched = len(results)

        sections = ""
        for url, live, saved, mismatches, error in results:
            all_match = not mismatches and not error
            icon = "✅" if all_match else "❌"
            hdr_bg = "#1e3a1e" if all_match else "#3a1a1a"

            if error:
                body_rows = f'<tr><td colspan="3" style="padding:10px 12px;color:{RED_FG};background:{RED_BG}">Error: {error}</td></tr>'
            else:
                body_rows = ""
                for key, label in FIELDS:
                    lv = live.get(key, "") or "—"
                    sv = saved.get(key, "") or "(not set)"
                    match = key not in mismatches
                    cell_bg = GREEN_BG if match else RED_BG
                    cell_fg = GREEN_FG if match else RED_FG
                    body_rows += f"""
                    <tr>
                      <td style="padding:8px 12px;font-weight:600;color:#555;white-space:nowrap;vertical-align:top;width:140px">{label}</td>
                      <td style="padding:8px 12px;color:#222;vertical-align:top">{lv}</td>
                      <td style="padding:8px 12px;background:{cell_bg};color:{cell_fg};vertical-align:top">{sv}</td>
                    </tr>"""

            sections += f"""
            <div style="margin-bottom:20px;border:1px solid #e5e5e5;border-radius:6px;overflow:hidden">
              <div style="background:{hdr_bg};padding:10px 16px">
                <span style="color:#fff;font-weight:bold;font-size:13px">{icon} {url}</span>
              </div>
              <table width="100%" cellpadding="0" cellspacing="0" style="font-size:13px;border-collapse:collapse">
                <tr style="background:#f9f9f9">
                  <th style="padding:8px 12px;text-align:left;color:#888;font-size:11px;text-transform:uppercase;width:140px">Field</th>
                  <th style="padding:8px 12px;text-align:left;color:#888;font-size:11px;text-transform:uppercase">Live</th>
                  <th style="padding:8px 12px;text-align:left;color:#888;font-size:11px;text-transform:uppercase">Saved</th>
                </tr>
                {body_rows}
              </table>
            </div>"""

        return f"""<!DOCTYPE html>
<html>
<body style="margin:0;padding:0;background:#f4f4f4;font-family:Arial,sans-serif">
  <table width="100%" cellpadding="0" cellspacing="0" style="padding:32px 0">
    <tr><td align="center">
      <table width="640" cellpadding="0" cellspacing="0" style="background:#fff;border-radius:8px;overflow:hidden;box-shadow:0 2px 8px rgba(0,0,0,0.08)">
        <tr>
          <td style="background:#1a1a1a;padding:28px 32px">
            <p style="margin:0;font-size:11px;color:#aaa;text-transform:uppercase;letter-spacing:1px">Site Monitor — Bulk Check</p>
            <h1 style="margin:6px 0 0;font-size:22px;color:#fff">{matched}/{total} URLs matched</h1>
          </td>
        </tr>
        <tr>
          <td style="padding:20px 32px 4px">
            <p style="margin:0;color:#444;font-size:14px">Checked at {now}. Green = matches your saved copy. Red = does not match.</p>
          </td>
        </tr>
        <tr><td style="padding:16px 32px 28px">{sections}</td></tr>
        <tr>
          <td style="padding:16px 32px;background:#f9f9f9;border-top:1px solid #eee">
            <p style="margin:0;font-size:12px;color:#aaa">You will only be notified for URLs where the live copy matches your saved copy. Green = matches. Red = does not match.</p>
          </td>
        </tr>
      </table>
    </td></tr>
  </table>
</body>
</html>"""

    # ── Resize ────────────────────────────────────────────────────────────────

    def _on_resize(self, _):
        self.update_idletasks()
        w = self._scroll.canvas.winfo_width()
        wrap = max(100, (w - COL_FIELD_W - 80) // 2)
        for lbl in self._result_labels:
            try:
                lbl.config(wraplength=wrap)
            except tk.TclError:
                pass


if __name__ == "__main__":
    App().mainloop()
