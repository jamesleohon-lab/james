import csv
import json
import subprocess
import sys
import threading
import tkinter as tk
from pathlib import Path
from tkinter import filedialog, ttk

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

CONFIG_FILE  = Path(__file__).parent / "config.json"
MONITOR_PY   = Path(__file__).parent / "monitor.py"
CHECKER_PY   = Path(__file__).parent / "app.py"
TASK_NAME    = "SiteMonitor-Pandora"
PYTHONW      = Path(sys.executable).parent / "pythonw.exe"

BG      = "#f5f5f5"
WHITE   = "#ffffff"
DARK    = "#1a1a1a"
MUTED   = "#888888"
BLUE    = "#2563eb"
GREEN   = "#16a34a"
DIVIDER = "#e0e0e0"

INTERVALS = [
    ("Every 15 minutes", "MINUTE", 15),
    ("Every 30 minutes", "MINUTE", 30),
    ("Every hour",       "HOURLY", 1),
    ("Every 2 hours",    "HOURLY", 2),
    ("Every 6 hours",    "HOURLY", 6),
    ("Every 12 hours",   "HOURLY", 12),
    ("Once a day",       "DAILY",  1),
]


# ── Helpers ───────────────────────────────────────────────────────────────────

def read_csv(path: str) -> list[dict]:
    with open(path, newline="", encoding="utf-8-sig") as f:
        reader = csv.DictReader(f)
        reader.fieldnames = [h.strip() for h in (reader.fieldnames or [])]
        return list(reader)

def read_xlsx(path: str) -> list[dict]:
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb.active
    rows = list(ws.iter_rows(values_only=True))
    if not rows:
        return []
    headers = [str(h).strip() if h is not None else "" for h in rows[0]]
    result = []
    for row in rows[1:]:
        if any(c for c in row if c is not None):
            result.append({headers[i]: (str(v).strip() if v is not None else "")
                           for i, v in enumerate(row)})
    return result

def load_cfg() -> dict:
    try:
        return json.loads(CONFIG_FILE.read_text(encoding="utf-8"))
    except Exception:
        return {}

def save_cfg(cfg: dict):
    CONFIG_FILE.write_text(json.dumps(cfg, indent=2, ensure_ascii=False), encoding="utf-8")

def task_status() -> str:
    """Returns 'enabled', 'disabled', or 'none'."""
    r = subprocess.run(
        ["schtasks", "/query", "/tn", TASK_NAME, "/fo", "LIST"],
        capture_output=True, text=True,
    )
    if r.returncode != 0:
        return "none"
    return "disabled" if "Disabled" in r.stdout else "enabled"

def create_or_update_task(schedule_type: str, modifier: int) -> str:
    exe = str(PYTHONW) if PYTHONW.exists() else sys.executable
    cmd = f'"{exe}" "{MONITOR_PY}"'
    r = subprocess.run([
        "schtasks", "/create", "/tn", TASK_NAME,
        "/tr", cmd,
        "/sc", schedule_type,
        "/mo", str(modifier),
        "/f",
    ], capture_output=True, text=True)
    if r.returncode != 0:
        return r.stderr.strip() or r.stdout.strip()
    # Explicitly enable after create in case a disabled task already existed
    enable_task()
    return ""

def enable_task():
    subprocess.run(["schtasks", "/change", "/tn", TASK_NAME, "/enable"], capture_output=True)

def disable_task():
    subprocess.run(["schtasks", "/change", "/tn", TASK_NAME, "/disable"], capture_output=True)

def delete_task():
    subprocess.run(["schtasks", "/delete", "/tn", TASK_NAME, "/f"], capture_output=True)


# ── Reusable widgets ──────────────────────────────────────────────────────────

def page_header(parent, title: str, subtitle: str = "", back_cmd=None) -> tk.Frame:
    hdr = tk.Frame(parent, bg=DARK, padx=28, pady=20)
    hdr.pack(fill="x")
    hdr.columnconfigure(1, weight=1)

    if back_cmd:
        back = tk.Button(hdr, text="← Back", font=("Arial", 10), bg=DARK, fg="#aaa",
                         relief="flat", cursor="hand2", activebackground=DARK,
                         activeforeground=WHITE, command=back_cmd)
        back.grid(row=0, column=0, rowspan=2, sticky="w", padx=(0, 20))

    col = 1 if back_cmd else 0
    if subtitle:
        tk.Label(hdr, text=subtitle, font=("Arial", 10), bg=DARK, fg=MUTED).grid(
            row=0, column=col, sticky="w")
    tk.Label(hdr, text=title, font=("Arial", 13, "bold"), bg=DARK, fg=WHITE).grid(
        row=1 if subtitle else 0, column=col, sticky="w", pady=(2, 0) if subtitle else 0)

    return hdr

def field_row(parent, row: int, label: str, var: tk.StringVar,
              show: str = "", width: int = 44) -> tk.Entry:
    tk.Label(parent, text=label, font=("Arial", 10, "bold"), bg=WHITE,
             fg="#444", anchor="w").grid(row=row, column=0, sticky="w", pady=(12, 2))
    entry = tk.Entry(parent, textvariable=var, font=("Arial", 11), show=show,
                     width=width, relief="flat", bg="#f9f9f9", fg=DARK,
                     insertbackground=DARK, highlightthickness=1,
                     highlightbackground=DIVIDER, highlightcolor=BLUE)
    entry.grid(row=row+1, column=0, sticky="ew", ipady=7)
    return entry

def section_label(parent, text: str, row: int):
    tk.Frame(parent, bg=DIVIDER, height=1).grid(row=row, column=0, sticky="ew", pady=(20, 0))
    tk.Label(parent, text=text.upper(), font=("Arial", 9, "bold"), bg=WHITE,
             fg=MUTED).grid(row=row+1, column=0, sticky="w", pady=(8, 0))


# ── Home page ─────────────────────────────────────────────────────────────────

class HomePage(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=BG)
        self._app = app
        self._build()

    def _build(self):
        page_header(self, "Site Monitor", "Welcome")

        body = tk.Frame(self, bg=BG, padx=32, pady=32)
        body.pack(fill="both", expand=True)

        tk.Label(body, text="What would you like to do?",
                 font=("Arial", 12), bg=BG, fg="#444").pack(anchor="w", pady=(0, 20))

        cards = tk.Frame(body, bg=BG)
        cards.pack(fill="x")
        cards.columnconfigure(0, weight=1)
        cards.columnconfigure(1, weight=1)

        self._card(cards, row=0, col=0,
                   icon="🔍",
                   title="SEO Copy Checker",
                   desc="Upload a spreadsheet of URLs and expected copy.\nRun checks manually and see live vs saved results\nside-by-side, colour-coded.",
                   btn_text="Open Checker →",
                   btn_color=BLUE,
                   cmd=self._open_checker)

        self._card(cards, row=0, col=1,
                   icon="🔔",
                   title="Email Notifications",
                   desc="Set a URL to watch automatically on a schedule.\nGet emailed when the live copy matches\nwhat you've saved.",
                   btn_text="Set Up →",
                   btn_color=GREEN,
                   cmd=lambda: self._app.show("notifications"))

        self._card(cards, row=1, col=0,
                   icon="↪️",
                   title="Redirect Tracker",
                   desc="Bulk upload a list of URLs and see the full\nredirect chain for each one — every hop,\nstatus code, and final destination.",
                   btn_text="Open →",
                   btn_color="#7c3aed",
                   cmd=lambda: self._app.show("redirects"))

        self._card(cards, row=1, col=1,
                   icon="🕷️",
                   title="SEO Crawler",
                   desc="Upload a list of URLs and crawl them all.\nExtracts live page titles, meta descriptions,\nand intro copy — export results to Excel.",
                   btn_text="Open →",
                   btn_color="#0d9488",
                   cmd=lambda: self._app.show("crawler"))

    def _card(self, parent, row, col, icon, title, desc, btn_text, btn_color, cmd):
        pad_x = (0, 8) if col == 0 else (8, 0)
        pad_y = (0, 8) if row == 0 else (8, 0)
        card = tk.Frame(parent, bg=WHITE, relief="flat",
                        highlightthickness=1, highlightbackground=DIVIDER)
        card.grid(row=row, column=col, sticky="nsew", padx=pad_x, pady=pad_y)

        inner = tk.Frame(card, bg=WHITE, padx=28, pady=28)
        inner.pack(fill="both", expand=True)

        tk.Label(inner, text=icon, font=("Arial", 28), bg=WHITE).pack(anchor="w")
        tk.Label(inner, text=title, font=("Arial", 13, "bold"), bg=WHITE,
                 fg=DARK).pack(anchor="w", pady=(8, 0))
        tk.Label(inner, text=desc, font=("Arial", 10), bg=WHITE, fg="#555",
                 justify="left", wraplength=260).pack(anchor="w", pady=(8, 20))
        tk.Button(inner, text=btn_text, font=("Arial", 10, "bold"),
                  bg=btn_color, fg=WHITE, relief="flat", padx=16, pady=8,
                  cursor="hand2",
                  activebackground=btn_color, activeforeground=WHITE,
                  command=cmd).pack(anchor="w")

        # Hover
        for w in [card, inner] + inner.winfo_children():
            w.bind("<Enter>", lambda _, c=card: c.config(highlightbackground=btn_color))
            w.bind("<Leave>", lambda _, c=card: c.config(highlightbackground=DIVIDER))

    def _open_checker(self):
        exe = str(PYTHONW) if PYTHONW.exists() else sys.executable
        subprocess.Popen([exe, str(CHECKER_PY)])


# ── Notifications page ────────────────────────────────────────────────────────

class NotificationsPage(tk.Frame):
    def __init__(self, parent, app):
        super().__init__(parent, bg=WHITE)
        self._app = app
        self._interval_var = tk.StringVar()
        self._status_var = tk.StringVar()
        self._mode = "single"           # "single" or "bulk"
        self._bulk_rows: list[dict] = []
        self._vars: dict[str, tk.StringVar] = {}
        self._build()

    def on_show(self):
        self._load_from_config()
        self._refresh_status()

    # ── Build ─────────────────────────────────────────────────────────────────

    def _build(self):
        page_header(self, "Email Notifications", "Site Monitor",
                    back_cmd=lambda: self._app.show("home"))

        canvas = tk.Canvas(self, bg=WHITE, highlightthickness=0)
        sb = tk.Scrollbar(self, orient="vertical", command=canvas.yview)
        canvas.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        canvas.pack(side="left", fill="both", expand=True)
        canvas.bind_all("<MouseWheel>",
                        lambda e: canvas.yview_scroll(-1*(e.delta//120), "units"))

        form_outer = tk.Frame(canvas, bg=WHITE)
        win = canvas.create_window((0, 0), window=form_outer, anchor="nw")
        form_outer.bind("<Configure>",
                        lambda _: canvas.configure(scrollregion=canvas.bbox("all")))
        canvas.bind("<Configure>", lambda e: canvas.itemconfig(win, width=e.width))

        form = tk.Frame(form_outer, bg=WHITE, padx=36, pady=24)
        form.pack(fill="both", expand=True)
        form.columnconfigure(0, weight=1)

        # ── Mode toggle ───────────────────────────────────────────────────────
        tk.Label(form, text="MONITORING MODE", font=("Arial", 9, "bold"),
                 bg=WHITE, fg=MUTED).grid(row=0, column=0, sticky="w")

        toggle = tk.Frame(form, bg=WHITE)
        toggle.grid(row=1, column=0, sticky="w", pady=(8, 0))

        self._single_btn = tk.Button(toggle, text="Single URL",
                                     font=("Arial", 10, "bold"), relief="flat",
                                     padx=18, pady=7, cursor="hand2",
                                     command=lambda: self._set_mode("single"))
        self._single_btn.pack(side="left")

        self._bulk_btn = tk.Button(toggle, text="Bulk Upload",
                                   font=("Arial", 10, "bold"), relief="flat",
                                   padx=18, pady=7, cursor="hand2",
                                   command=lambda: self._set_mode("bulk"))
        self._bulk_btn.pack(side="left", padx=(4, 0))

        # ── Single section ────────────────────────────────────────────────────
        self._single_section = tk.Frame(form, bg=WHITE)
        self._single_section.columnconfigure(0, weight=1)

        tk.Frame(self._single_section, bg=DIVIDER, height=1).grid(
            row=0, column=0, sticky="ew", pady=(20, 0))
        tk.Label(self._single_section, text="URL TO MONITOR", font=("Arial", 9, "bold"),
                 bg=WHITE, fg=MUTED).grid(row=1, column=0, sticky="w", pady=(8, 0))

        self._vars["url"] = tk.StringVar()
        field_row(self._single_section, 2, "URL", self._vars["url"])

        tk.Frame(self._single_section, bg=DIVIDER, height=1).grid(
            row=4, column=0, sticky="ew", pady=(20, 0))
        tk.Label(self._single_section, text="SAVED COPY TO COMPARE AGAINST",
                 font=("Arial", 9, "bold"), bg=WHITE, fg=MUTED).grid(
            row=5, column=0, sticky="w", pady=(8, 0))
        tk.Label(self._single_section,
                 text="A notification is sent when the live copy matches what you enter below.",
                 font=("Arial", 9), bg=WHITE, fg=MUTED).grid(
            row=6, column=0, sticky="w", pady=(2, 0))

        self._vars["title"] = tk.StringVar()
        field_row(self._single_section, 7, "Expected title", self._vars["title"])

        self._vars["meta_description"] = tk.StringVar()
        field_row(self._single_section, 9, "Expected meta description", self._vars["meta_description"])

        self._vars["intro_copy"] = tk.StringVar()
        field_row(self._single_section, 11, "Expected intro copy", self._vars["intro_copy"])

        self._vars["h1"] = tk.StringVar()
        field_row(self._single_section, 13, "Expected H1", self._vars["h1"])

        # ── Bulk section ──────────────────────────────────────────────────────
        self._bulk_section = tk.Frame(form, bg=WHITE)
        self._bulk_section.columnconfigure(0, weight=1)

        tk.Frame(self._bulk_section, bg=DIVIDER, height=1).grid(
            row=0, column=0, sticky="ew", pady=(20, 0))
        tk.Label(self._bulk_section, text="BULK UPLOAD", font=("Arial", 9, "bold"),
                 bg=WHITE, fg=MUTED).grid(row=1, column=0, sticky="w", pady=(8, 0))
        tk.Label(self._bulk_section,
                 text="Upload a CSV or XLSX file. Format: Col A = URL  |  Col B = Title  |  Col C = Meta Description  |  Col D = Intro Copy  |  Col E = H1",
                 font=("Arial", 9), bg=WHITE, fg=MUTED).grid(
            row=2, column=0, sticky="w", pady=(2, 0))

        bulk_btns = tk.Frame(self._bulk_section, bg=WHITE)
        bulk_btns.grid(row=3, column=0, sticky="w", pady=(12, 0))

        tk.Button(bulk_btns, text="Upload File (CSV or XLSX)",
                  font=("Arial", 10, "bold"), bg=DARK, fg=WHITE,
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#333", activeforeground=WHITE,
                  command=self._load_bulk_file).pack(side="left", padx=(0, 10))

        self._bulk_status_var = tk.StringVar(value="No file loaded")
        tk.Label(bulk_btns, textvariable=self._bulk_status_var,
                 font=("Arial", 10), bg=WHITE, fg=MUTED).pack(side="left")

        # URL listbox
        list_frame = tk.Frame(self._bulk_section, bg=WHITE)
        list_frame.grid(row=4, column=0, sticky="ew", pady=(10, 0))
        list_frame.columnconfigure(0, weight=1)

        self._url_listbox = tk.Listbox(list_frame, font=("Arial", 10), height=6,
                                        relief="flat", bg="#f9f9f9", fg=DARK,
                                        selectbackground=BLUE, selectforeground=WHITE,
                                        highlightthickness=1, highlightbackground=DIVIDER)
        self._url_listbox.grid(row=0, column=0, sticky="ew")
        list_sb = tk.Scrollbar(list_frame, orient="vertical",
                               command=self._url_listbox.yview)
        list_sb.grid(row=0, column=1, sticky="ns")
        self._url_listbox.configure(yscrollcommand=list_sb.set)

        # ── Common: schedule + email ──────────────────────────────────────────
        self._common_section = tk.Frame(form, bg=WHITE)
        self._common_section.columnconfigure(0, weight=1)

        tk.Frame(self._common_section, bg=DIVIDER, height=1).grid(
            row=0, column=0, sticky="ew", pady=(20, 0))
        tk.Label(self._common_section, text="CHECK FREQUENCY",
                 font=("Arial", 9, "bold"), bg=WHITE, fg=MUTED).grid(
            row=1, column=0, sticky="w", pady=(8, 0))

        interval_labels = [i[0] for i in INTERVALS]
        self._interval_var.set(interval_labels[2])
        tk.Label(self._common_section, text="How often to check",
                 font=("Arial", 10, "bold"), bg=WHITE, fg="#444", anchor="w").grid(
            row=2, column=0, sticky="w", pady=(10, 2))
        ttk.Combobox(self._common_section, textvariable=self._interval_var,
                     values=interval_labels, state="readonly",
                     font=("Arial", 11), width=30).grid(row=3, column=0, sticky="w", ipady=4)

        tk.Frame(self._common_section, bg=DIVIDER, height=1).grid(
            row=4, column=0, sticky="ew", pady=(20, 0))
        tk.Label(self._common_section, text="EMAIL SETTINGS",
                 font=("Arial", 9, "bold"), bg=WHITE, fg=MUTED).grid(
            row=5, column=0, sticky="w", pady=(8, 0))

        self._vars["gmail_sender"] = tk.StringVar()
        field_row(self._common_section, 6, "Gmail address (sender)", self._vars["gmail_sender"])

        self._vars["gmail_app_password"] = tk.StringVar()
        field_row(self._common_section, 8, "Gmail app password",
                  self._vars["gmail_app_password"], show="•")

        self._vars["notify_email"] = tk.StringVar()
        field_row(self._common_section, 10, "Send notifications to", self._vars["notify_email"])

        tk.Label(self._common_section,
                 text="Need an app password? Go to myaccount.google.com/apppasswords (2FA must be on).",
                 font=("Arial", 9), bg=WHITE, fg=MUTED).grid(
            row=12, column=0, sticky="w", pady=(6, 0))

        test_row = tk.Frame(self._common_section, bg=WHITE)
        test_row.grid(row=13, column=0, sticky="w", pady=(12, 0))

        tk.Button(test_row, text="Send Test Email",
                  font=("Arial", 10), bg="#f3f3f3", fg="#333",
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#e0e0e0",
                  command=self._send_test_email).pack(side="left", padx=(0, 10))

        self._test_result_var = tk.StringVar()
        tk.Label(test_row, textvariable=self._test_result_var,
                 font=("Arial", 10), bg=WHITE, fg=MUTED).pack(side="left")

        tk.Frame(self._common_section, bg=DIVIDER, height=1).grid(
            row=14, column=0, sticky="ew", pady=(20, 0))

        self._status_badge = tk.Label(self._common_section, textvariable=self._status_var,
                                      font=("Arial", 10, "bold"), bg=WHITE)
        self._status_badge.grid(row=15, column=0, sticky="w", pady=(14, 0))

        btn_row = tk.Frame(self._common_section, bg=WHITE)
        btn_row.grid(row=16, column=0, sticky="w", pady=(12, 32))

        tk.Button(btn_row, text="Save & Enable Schedule",
                  font=("Arial", 11, "bold"), bg=GREEN, fg=WHITE,
                  relief="flat", padx=16, pady=9, cursor="hand2",
                  activebackground="#15803d", activeforeground=WHITE,
                  command=self._save_and_enable).pack(side="left", padx=(0, 10))

        tk.Button(btn_row, text="Disable Schedule",
                  font=("Arial", 11), bg="#f3f3f3", fg="#555",
                  relief="flat", padx=16, pady=9, cursor="hand2",
                  activebackground="#e0e0e0",
                  command=self._disable).pack(side="left", padx=(0, 10))

        tk.Button(btn_row, text="Remove Schedule",
                  font=("Arial", 11), bg="#f3f3f3", fg="#c0392b",
                  relief="flat", padx=16, pady=9, cursor="hand2",
                  activebackground="#fdecea",
                  command=self._remove).pack(side="left")

        # Place common section in grid (always visible)
        self._common_section.grid(row=4, column=0, sticky="ew", in_=form)

        # Initial mode render
        self._set_mode("single")

    # ── Mode switching ────────────────────────────────────────────────────────

    def _set_mode(self, mode: str):
        self._mode = mode
        if mode == "single":
            self._single_btn.config(bg=DARK, fg=WHITE)
            self._bulk_btn.config(bg="#e8e8e8", fg="#555")
            self._bulk_section.grid_forget()
            self._single_section.grid(row=2, column=0, sticky="ew")
        else:
            self._bulk_btn.config(bg=DARK, fg=WHITE)
            self._single_btn.config(bg="#e8e8e8", fg="#555")
            self._single_section.grid_forget()
            self._bulk_section.grid(row=2, column=0, sticky="ew")

    # ── Bulk file loading ─────────────────────────────────────────────────────

    def _load_bulk_file(self):
        path = filedialog.askopenfilename(
            title="Select file",
            filetypes=[("Spreadsheet files", "*.csv *.xlsx"),
                       ("CSV", "*.csv"), ("Excel", "*.xlsx"), ("All", "*.*")],
        )
        if not path:
            return
        try:
            rows = read_xlsx(path) if path.lower().endswith(".xlsx") else read_csv(path)
        except Exception as e:
            self._bulk_status_var.set(f"Error reading file: {e}")
            return

        self._bulk_rows = [r for r in rows if r.get("URL", "").strip()]
        if not self._bulk_rows:
            self._bulk_status_var.set("No URLs found — check column A is labelled 'URL'.")
            return

        self._url_listbox.delete(0, tk.END)
        for r in self._bulk_rows:
            self._url_listbox.insert(tk.END, r["URL"].strip())

        count = len(self._bulk_rows)
        self._bulk_status_var.set(f"{count} URL{'s' if count > 1 else ''} loaded")

    # ── Config load / save ────────────────────────────────────────────────────

    def _load_from_config(self):
        cfg = load_cfg()
        mode = cfg.get("mode", "single")
        self._set_mode(mode)

        self._vars["gmail_sender"].set(cfg.get("gmail_sender", ""))
        self._vars["gmail_app_password"].set(cfg.get("gmail_app_password", ""))
        self._vars["notify_email"].set(cfg.get("notify_email", ""))

        # Single fields
        self._vars["url"].set(cfg.get("url", ""))
        saved = cfg.get("saved_copy", {})
        for k in ("title", "meta_description", "intro_copy", "h1"):
            v = saved.get(k, "")
            self._vars[k].set("" if v.startswith("PASTE THE EXACT") else v)

        # Bulk — repopulate listbox if bulk_urls exist in config
        bulk_urls = cfg.get("bulk_urls", [])
        if bulk_urls:
            self._bulk_rows = bulk_urls
            self._url_listbox.delete(0, tk.END)
            for r in bulk_urls:
                self._url_listbox.insert(tk.END, r.get("url", ""))
            count = len(bulk_urls)
            self._bulk_status_var.set(f"{count} URL{'s' if count > 1 else ''} loaded")

    def _collect_to_config(self) -> dict:
        cfg = load_cfg()
        cfg["mode"] = self._mode
        cfg["gmail_sender"] = self._vars["gmail_sender"].get().strip()
        cfg["gmail_app_password"] = self._vars["gmail_app_password"].get().strip()
        cfg["notify_email"] = self._vars["notify_email"].get().strip()

        if self._mode == "single":
            cfg["url"] = self._vars["url"].get().strip()
            cfg["saved_copy"] = {
                "title": self._vars["title"].get().strip(),
                "meta_description": self._vars["meta_description"].get().strip(),
                "intro_copy": self._vars["intro_copy"].get().strip(),
                "h1": self._vars["h1"].get().strip(),
            }
        else:
            cfg["bulk_urls"] = [
                {
                    "url": r.get("URL", r.get("url", "")).strip(),
                    "title": r.get("Title", r.get("title", "")).strip(),
                    "meta_description": r.get("Meta Description",
                                              r.get("meta_description", "")).strip(),
                    "intro_copy": r.get("Intro Copy", r.get("intro_copy", "")).strip(),
                    "h1": r.get("H1", r.get("h1", "")).strip(),
                }
                for r in self._bulk_rows
            ]
        return cfg

    # ── Test email ────────────────────────────────────────────────────────────

    def _send_test_email(self):
        import smtplib
        from email.mime.multipart import MIMEMultipart
        from email.mime.text import MIMEText

        sender   = self._vars["gmail_sender"].get().strip()
        password = self._vars["gmail_app_password"].get().strip()
        to       = self._vars["notify_email"].get().strip()

        if not sender or not password or not to:
            self._test_result_var.set("⚠️ Fill in all email fields first.")
            return

        self._test_result_var.set("Sending…")
        self.update_idletasks()

        html = """<!DOCTYPE html>
<html>
<body style="margin:0;padding:0;background:#f4f4f4;font-family:Arial,sans-serif">
  <table width="100%" cellpadding="0" cellspacing="0" style="padding:32px 0">
    <tr><td align="center">
      <table width="520" cellpadding="0" cellspacing="0"
             style="background:#fff;border-radius:8px;overflow:hidden;box-shadow:0 2px 8px rgba(0,0,0,0.08)">
        <tr>
          <td style="background:#1a1a1a;padding:28px 32px">
            <p style="margin:0;font-size:11px;color:#aaa;text-transform:uppercase;letter-spacing:1px">Site Monitor</p>
            <h1 style="margin:6px 0 0;font-size:22px;color:#fff">✅ Test email</h1>
          </td>
        </tr>
        <tr>
          <td style="padding:28px 32px">
            <p style="margin:0;color:#444;font-size:14px">
              Your email notifications are set up correctly.<br><br>
              You'll receive an email like this whenever monitored copy goes live.
            </p>
          </td>
        </tr>
      </table>
    </td></tr>
  </table>
</body>
</html>"""

        try:
            msg = MIMEMultipart("alternative")
            msg["Subject"] = "[Site Monitor] Test email"
            msg["From"]    = sender
            msg["To"]      = to
            msg.attach(MIMEText(html, "html"))
            with smtplib.SMTP_SSL("smtp.gmail.com", 465) as server:
                server.login(sender, password)
                server.sendmail(sender, to, msg.as_string())
            self._test_result_var.set(f"✅ Sent to {to}")
        except Exception as e:
            self._test_result_var.set(f"❌ Failed: {e}")

    # ── Schedule actions ──────────────────────────────────────────────────────

    def _refresh_status(self):
        status = task_status()
        if status == "enabled":
            self._status_var.set("● Schedule is active")
            self._status_badge.config(fg="#16a34a")
        elif status == "disabled":
            self._status_var.set("○ Schedule is disabled")
            self._status_badge.config(fg="#d97706")
        else:
            self._status_var.set("○ No schedule set up yet")
            self._status_badge.config(fg=MUTED)

    def _save_and_enable(self):
        if self._mode == "bulk" and not self._bulk_rows:
            self._status_var.set("Please upload a file first.")
            self._status_badge.config(fg="#b71c1c")
            return
        cfg = self._collect_to_config()
        save_cfg(cfg)
        label = self._interval_var.get()
        _, sc_type, modifier = next(i for i in INTERVALS if i[0] == label)
        err = create_or_update_task(sc_type, modifier)
        if err:
            self._status_var.set(f"Error creating task: {err}")
            self._status_badge.config(fg="#b71c1c")
            return
        self._refresh_status()

    def _disable(self):
        disable_task()
        self._refresh_status()

    def _remove(self):
        delete_task()
        self._refresh_status()


# ── Redirect tracker page ─────────────────────────────────────────────────────

class RedirectTrackerPage(tk.Frame):
    PURPLE = "#7c3aed"

    def __init__(self, parent, app):
        super().__init__(parent, bg=WHITE)
        self._app = app
        self._urls: list[str] = []
        self._build()

    def _build(self):
        from redirect_checker import get_redirect_chain, status_colour

        page_header(self, "Redirect Tracker", "Site Monitor",
                    back_cmd=lambda: self._app.show("home"))

        # ── Top bar ───────────────────────────────────────────────────────────
        bar = tk.Frame(self, bg="#ececec", padx=24, pady=12)
        bar.pack(fill="x")

        tk.Button(bar, text="Upload File (CSV or XLSX)",
                  font=("Arial", 10, "bold"), bg=DARK, fg=WHITE,
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#333", activeforeground=WHITE,
                  command=self._load_file).pack(side="left", padx=(0, 10))

        self._check_btn = tk.Button(bar, text="Check All",
                  font=("Arial", 10, "bold"), bg=self.PURPLE, fg=WHITE,
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#6d28d9", activeforeground=WHITE,
                  state="disabled", command=self._run_checks)
        self._check_btn.pack(side="left", padx=(0, 10))

        self._export_btn = tk.Button(bar, text="Export CSV",
                  font=("Arial", 10), bg="#ececec", fg="#555",
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#ddd", state="disabled",
                  command=self._export_csv)
        self._export_btn.pack(side="left")

        self._file_label = tk.Label(bar, text="No file loaded",
                                    font=("Arial", 10), bg="#ececec", fg=MUTED)
        self._file_label.pack(side="left", padx=(16, 0))

        tk.Label(bar,
                 text="Format: Col A = URL  (row 1 = header labelled 'URL')",
                 font=("Arial", 9), bg="#ececec", fg=MUTED).pack(
            side="right", padx=(0, 4))

        # ── Column headers ────────────────────────────────────────────────────
        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")
        col_hdr = tk.Frame(self, bg="#f0f0f0", padx=24, pady=8)
        col_hdr.pack(fill="x")
        tk.Label(col_hdr, text="ORIGINAL URL", font=("Arial", 9, "bold"),
                 bg="#f0f0f0", fg=MUTED, width=40, anchor="w").pack(side="left")
        tk.Label(col_hdr, text="REDIRECT CHAIN",
                 font=("Arial", 9, "bold"), bg="#f0f0f0", fg=MUTED, anchor="w").pack(side="left")
        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")

        # ── Scrollable results ────────────────────────────────────────────────
        wrap = tk.Frame(self, bg=WHITE)
        wrap.pack(fill="both", expand=True)

        self._canvas = tk.Canvas(wrap, bg=WHITE, highlightthickness=0)
        sb = tk.Scrollbar(wrap, orient="vertical", command=self._canvas.yview)
        self._inner = tk.Frame(self._canvas, bg=WHITE)
        self._canvas.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self._canvas.pack(side="left", fill="both", expand=True)
        self._win = self._canvas.create_window((0, 0), window=self._inner, anchor="nw")
        self._inner.bind("<Configure>",
                         lambda _: self._canvas.configure(
                             scrollregion=self._canvas.bbox("all")))
        self._canvas.bind("<Configure>",
                          lambda e: self._canvas.itemconfig(self._win, width=e.width))
        self._canvas.bind_all("<MouseWheel>",
                              lambda e: self._canvas.yview_scroll(
                                  -1*(e.delta//120), "units"))

        # ── Status bar ────────────────────────────────────────────────────────
        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")
        self._status_var = tk.StringVar(value="Upload a file and press Check All to begin.")
        tk.Frame(self, bg="#f0f0f0", padx=24, pady=10).pack(fill="x")
        self._status_lbl = tk.Label(self, textvariable=self._status_var,
                                    font=("Arial", 10), bg="#f0f0f0", fg=MUTED, anchor="w",
                                    padx=24)
        self._status_lbl.pack(fill="x")

        # Store results for export
        self._results: list[tuple[str, list]] = []

    # ── File loading ──────────────────────────────────────────────────────────

    def _load_file(self):
        path = filedialog.askopenfilename(
            title="Select file",
            filetypes=[("Spreadsheet files", "*.csv *.xlsx"),
                       ("CSV", "*.csv"), ("Excel", "*.xlsx"), ("All", "*.*")],
        )
        if not path:
            return
        try:
            rows = read_xlsx(path) if path.lower().endswith(".xlsx") else read_csv(path)
        except Exception as e:
            self._status_var.set(f"Failed to read file: {e}")
            return

        self._urls = [r.get("URL", "").strip() for r in rows if r.get("URL", "").strip()]
        if not self._urls:
            self._status_var.set("No URLs found — check column A is labelled 'URL'.")
            return

        count = len(self._urls)
        self._file_label.config(text=f"{count} URL{'s' if count > 1 else ''} loaded")
        self._check_btn.config(state="normal")
        self._status_var.set(f"{count} URL{'s' if count > 1 else ''} ready. Press Check All to start.")
        self._clear_results()

    # ── Checking ──────────────────────────────────────────────────────────────

    def _run_checks(self):
        self._clear_results()
        self._results.clear()
        self._check_btn.config(state="disabled", text="Checking…")
        self._export_btn.config(state="disabled")
        threading.Thread(target=self._do_checks, daemon=True).start()

    def _do_checks(self):
        from redirect_checker import get_redirect_chain
        total = len(self._urls)
        for i, url in enumerate(self._urls):
            self.after(0, self._status_var.set,
                       f"Checking {i+1}/{total}: {url}")
            chain = get_redirect_chain(url)
            self._results.append((url, chain))
            self.after(0, self._render_row, url, chain, i)

        def done():
            hops = sum(len(c) - 1 for _, c in self._results if len(c) > 1)
            redirected = sum(1 for _, c in self._results if len(c) > 1)
            self._status_var.set(
                f"Done — {total} URL{'s' if total > 1 else ''} checked, "
                f"{redirected} redirect{'s' if redirected != 1 else ''} found.")
            self._check_btn.config(state="normal", text="Check All")
            self._export_btn.config(state="normal")

        self.after(0, done)

    # ── Rendering ─────────────────────────────────────────────────────────────

    def _clear_results(self):
        for w in self._inner.winfo_children():
            w.destroy()

    def _render_row(self, url: str, chain: list, row_idx: int):
        from redirect_checker import status_colour

        row_bg = WHITE if row_idx % 2 == 0 else "#fafafa"
        row_frame = tk.Frame(self._inner, bg=row_bg, padx=24, pady=12)
        row_frame.pack(fill="x")

        # Original URL (fixed width column)
        url_lbl = tk.Label(row_frame, text=url, font=("Arial", 10),
                           bg=row_bg, fg="#222", width=38, anchor="nw",
                           wraplength=280, justify="left")
        url_lbl.pack(side="left", anchor="n", padx=(0, 16))

        # Chain column
        chain_frame = tk.Frame(row_frame, bg=row_bg)
        chain_frame.pack(side="left", fill="x", expand=True, anchor="n")

        for step_i, step in enumerate(chain):
            hop = tk.Frame(chain_frame, bg=row_bg)
            hop.pack(anchor="w", pady=(0, 4))

            # Arrow (not on first step)
            if step_i > 0:
                tk.Label(hop, text="↳", font=("Arial", 11),
                         bg=row_bg, fg=MUTED).pack(side="left", padx=(0, 6))

            # Status badge
            if step["status"] is not None or step.get("error"):
                sbg, sfg = status_colour(step["status"])
                badge_text = step["label"] if step["status"] else "Error"
                tk.Label(hop, text=badge_text, font=("Arial", 9, "bold"),
                         bg=sbg, fg=sfg, padx=6, pady=2).pack(side="left", padx=(0, 8))

            # URL for this step
            display_url = step["url"] if len(step["url"]) <= 70 else step["url"][:67] + "…"
            tk.Label(hop, text=display_url, font=("Arial", 10),
                     bg=row_bg, fg="#333").pack(side="left")

            # Error message if any
            if step.get("error"):
                tk.Label(hop, text=f"  ({step['error']})",
                         font=("Arial", 9), bg=row_bg, fg="#b71c1c").pack(side="left")

        tk.Frame(self._inner, bg=DIVIDER, height=1).pack(fill="x")

    # ── Export ────────────────────────────────────────────────────────────────

    def _export_csv(self):
        path = filedialog.asksaveasfilename(
            title="Save results",
            defaultextension=".csv",
            initialfile="redirect-results.csv",
            filetypes=[("CSV files", "*.csv")],
        )
        if not path:
            return
        with open(path, "w", newline="", encoding="utf-8") as f:
            writer = csv.writer(f)
            writer.writerow(["Original URL", "Hop", "URL", "Status", "Label", "Error"])
            for orig_url, chain in self._results:
                for i, step in enumerate(chain):
                    writer.writerow([
                        orig_url if i == 0 else "",
                        i + 1,
                        step["url"],
                        step.get("status", ""),
                        step.get("label", ""),
                        step.get("error", ""),
                    ])
        self._status_var.set(f"Exported to {path}")


# ── SEO Crawler page ──────────────────────────────────────────────────────────

class CrawlerPage(tk.Frame):
    TEAL = "#0d9488"

    def __init__(self, parent, app):
        super().__init__(parent, bg=WHITE)
        self._app = app
        self._urls: list[str] = []
        self._results: list[dict] = []
        self._build()

    def _build(self):
        from monitor import fetch_copy  # imported here to avoid circular at module level

        page_header(self, "SEO Crawler", "Site Monitor",
                    back_cmd=lambda: self._app.show("home"))

        # ── Top bar ───────────────────────────────────────────────────────────
        bar = tk.Frame(self, bg="#ececec", padx=24, pady=12)
        bar.pack(fill="x")

        tk.Button(bar, text="Upload URLs (CSV or XLSX)",
                  font=("Arial", 10, "bold"), bg=DARK, fg=WHITE,
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#333", activeforeground=WHITE,
                  command=self._load_file).pack(side="left", padx=(0, 10))

        self._crawl_btn = tk.Button(bar, text="Crawl All",
                  font=("Arial", 10, "bold"), bg=self.TEAL, fg=WHITE,
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#0f766e", activeforeground=WHITE,
                  state="disabled", command=self._run_crawl)
        self._crawl_btn.pack(side="left", padx=(0, 10))

        self._export_btn = tk.Button(bar, text="Export to Excel",
                  font=("Arial", 10), bg="#ececec", fg="#555",
                  relief="flat", padx=12, pady=6, cursor="hand2",
                  activebackground="#ddd", state="disabled",
                  command=self._export_xlsx)
        self._export_btn.pack(side="left")

        self._file_label = tk.Label(bar, text="No file loaded",
                                    font=("Arial", 10), bg="#ececec", fg=MUTED)
        self._file_label.pack(side="left", padx=(16, 0))

        tk.Label(bar, text="Format: Col A = URL  (row 1 = header labelled 'URL')",
                 font=("Arial", 9), bg="#ececec", fg=MUTED).pack(side="right", padx=(0, 4))

        # ── Column headers ────────────────────────────────────────────────────
        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")
        col_hdr = tk.Frame(self, bg="#f0f0f0", padx=24, pady=8)
        col_hdr.pack(fill="x")
        col_hdr.columnconfigure(0, weight=2)
        col_hdr.columnconfigure(1, weight=2)
        col_hdr.columnconfigure(2, weight=2)
        col_hdr.columnconfigure(3, weight=2)
        col_hdr.columnconfigure(4, weight=3)
        for ci, txt in enumerate(["URL", "H1", "TITLE", "META DESCRIPTION", "INTRO COPY"]):
            tk.Label(col_hdr, text=txt, anchor="w", font=("Arial", 9, "bold"),
                     bg="#f0f0f0", fg=MUTED).grid(row=0, column=ci, sticky="ew", padx=(0, 16))
        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")

        # ── Scrollable results ────────────────────────────────────────────────
        wrap = tk.Frame(self, bg=WHITE)
        wrap.pack(fill="both", expand=True)

        self._canvas = tk.Canvas(wrap, bg=WHITE, highlightthickness=0)
        sb = tk.Scrollbar(wrap, orient="vertical", command=self._canvas.yview)
        self._inner = tk.Frame(self._canvas, bg=WHITE)
        self._canvas.configure(yscrollcommand=sb.set)
        sb.pack(side="right", fill="y")
        self._canvas.pack(side="left", fill="both", expand=True)
        self._win = self._canvas.create_window((0, 0), window=self._inner, anchor="nw")
        self._inner.bind("<Configure>",
                         lambda _: self._canvas.configure(scrollregion=self._canvas.bbox("all")))
        self._canvas.bind("<Configure>",
                          lambda e: self._canvas.itemconfig(self._win, width=e.width))
        self._canvas.bind_all("<MouseWheel>",
                              lambda e: self._canvas.yview_scroll(-1*(e.delta//120), "units"))

        # ── Status bar ────────────────────────────────────────────────────────
        tk.Frame(self, bg=DIVIDER, height=1).pack(fill="x")
        self._status_var = tk.StringVar(value="Upload a file and press Crawl All to begin.")
        status_frame = tk.Frame(self, bg="#f0f0f0", padx=24, pady=10)
        status_frame.pack(fill="x")
        tk.Label(status_frame, textvariable=self._status_var, font=("Arial", 10),
                 bg="#f0f0f0", fg=MUTED, anchor="w").pack(fill="x")

    # ── File loading ──────────────────────────────────────────────────────────

    def _load_file(self):
        path = filedialog.askopenfilename(
            title="Select file",
            filetypes=[("Spreadsheet files", "*.csv *.xlsx"),
                       ("CSV", "*.csv"), ("Excel", "*.xlsx"), ("All", "*.*")],
        )
        if not path:
            return
        try:
            rows = read_xlsx(path) if path.lower().endswith(".xlsx") else read_csv(path)
        except Exception as e:
            self._status_var.set(f"Failed to read file: {e}")
            return

        self._urls = [r.get("URL", "").strip() for r in rows if r.get("URL", "").strip()]
        if not self._urls:
            self._status_var.set("No URLs found — check column A is labelled 'URL'.")
            return

        count = len(self._urls)
        self._file_label.config(text=f"{count} URL{'s' if count > 1 else ''} loaded")
        self._crawl_btn.config(state="normal")
        self._status_var.set(f"{count} URL{'s' if count > 1 else ''} ready. Press Crawl All to start.")
        self._clear_results()
        self._results.clear()

    # ── Crawling ──────────────────────────────────────────────────────────────

    def _run_crawl(self):
        self._clear_results()
        self._results.clear()
        self._crawl_btn.config(state="disabled", text="Crawling…")
        self._export_btn.config(state="disabled")
        threading.Thread(target=self._do_crawl, daemon=True).start()

    def _do_crawl(self):
        from monitor import fetch_copy
        total = len(self._urls)
        for i, url in enumerate(self._urls):
            self.after(0, self._status_var.set, f"Crawling {i+1}/{total}: {url}")
            try:
                live = fetch_copy(url)
                result = {"url": url, "error": "", **live}
            except Exception as e:
                result = {"url": url, "title": "", "meta_description": "",
                          "intro_copy": "", "error": str(e)}
            self._results.append(result)
            self.after(0, self._render_row, result, i)

        def done():
            errors = sum(1 for r in self._results if r["error"])
            msg = f"Done — {total} URL{'s' if total > 1 else ''} crawled"
            if errors:
                msg += f", {errors} error{'s' if errors != 1 else ''}"
            self._status_var.set(msg + ".")
            self._crawl_btn.config(state="normal", text="Crawl All")
            self._export_btn.config(state="normal")

        self.after(0, done)

    # ── Rendering ─────────────────────────────────────────────────────────────

    def _clear_results(self):
        for w in self._inner.winfo_children():
            w.destroy()

    def _render_row(self, result: dict, row_idx: int):
        row_bg = WHITE if row_idx % 2 == 0 else "#fafafa"
        row_frame = tk.Frame(self._inner, bg=row_bg, padx=24, pady=10)
        row_frame.pack(fill="x")
        row_frame.columnconfigure(0, weight=2)
        row_frame.columnconfigure(1, weight=2)
        row_frame.columnconfigure(2, weight=2)
        row_frame.columnconfigure(3, weight=2)
        row_frame.columnconfigure(4, weight=3)

        if result["error"]:
            tk.Label(row_frame, text=result["url"], font=("Arial", 10, "bold"),
                     bg=row_bg, fg="#222", anchor="nw", wraplength=160,
                     justify="left").grid(row=0, column=0, sticky="nsew", padx=(0, 12))
            tk.Label(row_frame, text=f"Error: {result['error']}", font=("Arial", 10),
                     bg="#fdecea", fg="#b71c1c", anchor="nw", wraplength=400,
                     justify="left", padx=8, pady=4).grid(
                row=0, column=1, columnspan=4, sticky="nsew")
        else:
            fields = [
                ("url",              160),
                ("h1",               160),
                ("title",            160),
                ("meta_description", 180),
                ("intro_copy",       200),
            ]
            for ci, (key, wrap) in enumerate(fields):
                val = result.get(key, "") or "—"
                tk.Label(row_frame, text=val, font=("Arial", 10),
                         bg=row_bg, fg="#222", anchor="nw",
                         wraplength=wrap, justify="left").grid(
                    row=0, column=ci, sticky="nsew", padx=(0, 12))

        tk.Frame(self._inner, bg=DIVIDER, height=1).pack(fill="x")

    # ── Export ────────────────────────────────────────────────────────────────

    def _export_xlsx(self):
        path = filedialog.asksaveasfilename(
            title="Save crawl results",
            defaultextension=".xlsx",
            initialfile="crawl-results.xlsx",
            filetypes=[("Excel files", "*.xlsx")],
        )
        if not path:
            return

        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Crawl Results"

        headers = ["URL", "H1", "Title", "Meta Description", "Intro Copy", "Error"]
        col_widths = [50, 45, 55, 70, 80, 40]
        header_fill = PatternFill("solid", fgColor="1A1A1A")
        header_font = Font(bold=True, color="FFFFFF", size=11)

        for ci, (h, w) in enumerate(zip(headers, col_widths), start=1):
            cell = ws.cell(row=1, column=ci, value=h)
            cell.font = header_font
            cell.fill = header_fill
            cell.alignment = Alignment(vertical="center")
            ws.column_dimensions[cell.column_letter].width = w
        ws.row_dimensions[1].height = 22

        for ri, result in enumerate(self._results, start=2):
            vals = [
                result.get("url", ""),
                result.get("h1", ""),
                result.get("title", ""),
                result.get("meta_description", ""),
                result.get("intro_copy", ""),
                result.get("error", ""),
            ]
            for ci, val in enumerate(vals, start=1):
                cell = ws.cell(row=ri, column=ci, value=val)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
            ws.row_dimensions[ri].height = 40

        wb.save(path)
        self._status_var.set(f"Exported to {path}")


# ── App shell ─────────────────────────────────────────────────────────────────

class SiteMonitorApp(tk.Tk):
    def __init__(self):
        super().__init__()
        self.title("Site Monitor")
        self.configure(bg=BG)
        self.minsize(720, 480)
        self.resizable(True, True)

        self._pages: dict[str, tk.Frame] = {}
        self._current: tk.Frame | None = None

        container = tk.Frame(self, bg=BG)
        container.pack(fill="both", expand=True)

        self._pages["home"]          = HomePage(container, self)
        self._pages["notifications"] = NotificationsPage(container, self)
        self._pages["redirects"]     = RedirectTrackerPage(container, self)
        self._pages["crawler"]       = CrawlerPage(container, self)

        self.show("home")
        self.geometry("900x660")
        self.update_idletasks()
        w, h = self.winfo_width(), self.winfo_height()
        sw, sh = self.winfo_screenwidth(), self.winfo_screenheight()
        self.geometry(f"{w}x{h}+{(sw-w)//2}+{(sh-h)//2}")

    def show(self, name: str):
        if self._current:
            self._current.pack_forget()
        page = self._pages[name]
        page.pack(fill="both", expand=True)
        if hasattr(page, "on_show"):
            page.on_show()
        self._current = page


if __name__ == "__main__":
    SiteMonitorApp().mainloop()
